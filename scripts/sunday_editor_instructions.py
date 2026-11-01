"""The standing IPR Editor Instructions, rendered for every editorial email.

``docs/IPR_EDITOR_INSTRUCTIONS.md`` is the single source. Only the editor-facing
part travels: everything from the owner-only note onward is stripped. This is a
presentation transform with no model, network or approval authority, and it
fails closed so an editor email is never sent without the instructions.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from scripts.sunday_editor_readable import ReadableEditorError

SOURCE = Path(__file__).resolve().parents[1] / "docs" / "IPR_EDITOR_INSTRUCTIONS.md"
FILENAME = "IPR-Editor-Instructions.docx"
HEADING = "# IPR Editor Instructions"
OWNER_SEPARATOR = "\n---\n\n### Note for the owner"
INLINE = re.compile(r"(\*\*[^*\n]+\*\*|\*[^*\n]+\*)")
BULLET = re.compile(r"^- (.*)$")
NUMBERED = re.compile(r"^([0-9]+\. .*)$")


class EditorInstructionsError(ReadableEditorError):
    """The standing instructions are missing or not in their reviewed shape."""


def editor_instructions_markdown(path=SOURCE):
    """Editor-facing Markdown only; refuse a missing or reshaped source."""
    try:
        text = Path(path).read_text(encoding="utf-8")
    except OSError as exc:
        raise EditorInstructionsError("editor instructions unreadable: {}".format(exc)) from exc
    if not text.startswith(HEADING + "\n") or text.count(OWNER_SEPARATOR) != 1:
        raise EditorInstructionsError("editor instructions heading or owner note boundary missing")
    editor_part = text.split(OWNER_SEPARATOR, 1)[0].strip()
    if "## 1. Who does what" not in editor_part:
        raise EditorInstructionsError("editor instructions body missing")
    return editor_part + "\n"


def _blocks(markdown):
    """``(kind, text)`` with wrapped lines joined; kinds follow the doc's own idiom.

    Numbered items keep their literal number: Word's shared List Number
    sequence would otherwise continue §3's list into §8.
    """
    blocks, current = [], None
    for line in markdown.splitlines():
        if not line.strip():
            current = None
            continue
        bullet, numbered = BULLET.match(line), NUMBERED.match(line)
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("h{}".format(level), line[level:].strip()))
            current = None
        elif bullet or numbered:
            current = ["bullet" if bullet else "number", (bullet or numbered).group(1)]
            blocks.append(current)
        elif current is not None:
            current[1] += " " + line.strip()
        else:
            current = ["para", line.strip()]
            blocks.append(current)
    return [tuple(block) for block in blocks]


def _runs(paragraph, text):
    for part in INLINE.split(text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part.replace("`", ""))


def render_instructions_docx(path=SOURCE):
    """Deterministic .docx bytes of the editor-facing instructions."""
    doc = Document()
    doc.core_properties.title = "IPR Editor Instructions"
    doc.core_properties.author = "Indo-Pacific Record"
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.color.rgb = None
    styles = {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2",
              "bullet": "List Bullet", "number": None, "para": None}
    for kind, text in _blocks(editor_instructions_markdown(path)):
        _runs(doc.add_paragraph(style=styles.get(kind, "Heading 2")), text)

    output = io.BytesIO()
    doc.save(output)
    # Same stable-ZIP rule as the manuscript copy: identical bytes in the owner
    # preview and the editor email, so both SHA lines can be compared.
    stable = io.BytesIO()
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(stable, "w") as target:
        for entry in source.infolist():
            entry.date_time = (1980, 1, 1, 0, 0, 0)
            target.writestr(entry, source.read(entry.filename))
    return stable.getvalue()
