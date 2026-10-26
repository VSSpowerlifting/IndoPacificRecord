"""Offline acceptance tests for the Sunday Briefs narrative-writing reform.

Nothing here calls a model, SMTP server, database writer or publisher: the only
client is an in-process fake, and every mail path is patched to fail the test.
All manuscript text is the synthetic Tavaran Navy fixture.
"""
from __future__ import annotations

import copy
import io
import json
import re
import tempfile
import unittest
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from core.brief_editorial_evidence import load_editorial_evidence
from core.brief_editorial_source_use import SourceUseError, summarize_source_use
from core.brief_prose_standard import PROMPT_STANDARD, check_article, fingerprint
from scripts import analyze_editor_return as analyzer
from scripts.import_sunday_editor_docx import reconstruct_return
from scripts.sunday_briefs_auto_writer import (
    compose, validate_manuscript, validate_prose_boundaries, writing_schema)
from scripts.sunday_editor_docx import _references, render_editor_docx
from scripts.sunday_editor_readable import (
    ReadableEditorError, editable_sections, prose_only_html, prose_only_text)
from scripts.sunday_editorial_handoff import render_packet
from scripts.validate_editorial_return import ReturnValidationError, validate_return
from tests.briefs_narrative_fixtures import (
    ARTICLE_BLOCKS, AUDIT, SAT, article_manuscript, evidence, legacy_manuscript, scaffold)

DESKS = {1: "china", 2: "singapore"}
PRIVATE_MARKERS = (
    AUDIT["baseline_and_routine"][:40], AUDIT["limitations"][:40],
    AUDIT["terms_to_verify"][:40], AUDIT["cross_desk_assessment"][:40],
    "NARRATIVE QUALITY", "ARTICLE SOURCE ACCOUNTING", "Block 1 | fp",
    "SOURCE RECORD IDS", "EXTERNAL SOURCE IDS", "Limitations ledger",
)


def report(manuscript):
    return check_article(manuscript, desk_by_id=DESKS)


def codes(result, kind="flags"):
    return {item["code"] for item in result[kind]}


def with_blocks(*blocks):
    """Article manuscript whose blocks are ``(heading, text)`` pairs citing record 2."""
    m = article_manuscript()
    m["article"] = [{"heading": h, "text": t, "citations": [2]} for h, t in blocks]
    return m


def fake_client(*outputs):
    calls = []
    queue = list(outputs)

    def stream(**kwargs):
        calls.append(kwargs)
        block = SimpleNamespace(type="tool_use", name="compose_editorial_draft",
                                input=queue.pop(0))
        response = SimpleNamespace(stop_reason="tool_use", content=[block])
        return nullcontext(SimpleNamespace(get_final_message=lambda: response))

    return SimpleNamespace(messages=SimpleNamespace(stream=stream)), calls


def run_compose(*outputs, **kwargs):
    client, calls = fake_client(*outputs)
    with patch("scripts.sunday_briefs_auto_writer.choose_evidence", return_value=evidence()), \
            patch("core.daily_machine_records.editorial_candidates", return_value=[]):
        result = compose({"week_start": "2026-10-04", "week_ending": SAT,
                          "desks": ["china", "singapore"], "source_trail": []},
                         SAT, client=client, **kwargs)
    return result, calls


def docx_text(data):
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    parts.extend(str(v) for v in (doc.core_properties.title, doc.core_properties.author,
                                  doc.core_properties.comments, doc.core_properties.subject))
    return "\n".join(parts)


def edit_docx(data, mutate):
    doc = Document(io.BytesIO(data))
    mutate(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def set_text(paragraph, value):
    for run in paragraph.runs[1:]:
        run.text = ""
    paragraph.runs[0].text = value


class NarrativeStandardTests(unittest.TestCase):
    """Acceptance checks 1, 2, 3, 5, 6 and 7."""

    def test_technical_development_reads_naturally_and_passes_every_check(self):
        result = report(article_manuscript())
        self.assertEqual((result["errors"], result["flags"], result["hold"]), ([], [], False))
        article = "\n".join(text for _, text in ARTICLE_BLOCKS)
        self.assertIn("A towed array is a long cable of underwater microphones", article)

    def test_opening_presents_the_event_before_the_source(self):
        lead = ARTICLE_BLOCKS[0][1]
        self.assertTrue(lead.startswith("The Tavaran Navy has spent three weeks"))
        bad = with_blocks(
            ("", "On 16 October 2026 the Tavaran Navy published a press release describing a "
                 "sonar trial. The release was posted to its website in the morning."),
            *ARTICLE_BLOCKS[1:])
        self.assertIn("lead_source_first", codes(report(bad)))
        self.assertIn("lead_source_first", codes(report(legacy_manuscript())))
        also_bad = with_blocks(
            ("", "According to the stored record, a navy statement discussed a new sonar system "
                 "that the navy said it had tested over several weeks in a strait."),
            *ARTICLE_BLOCKS[1:])
        self.assertIn("lead_source_first", codes(report(also_bad)))

    def test_introductory_material_is_not_restated_mechanically(self):
        lead = ARTICLE_BLOCKS[0][1]
        bad = with_blocks(*ARTICLE_BLOCKS[:3], ("Recap", lead))
        found = codes(report(bad), "flags") | codes(report(bad), "errors")
        self.assertTrue({"restated_lead", "duplicate_paragraph"} & found)
        stock = with_blocks(("", ARTICLE_BLOCKS[0][1]), ("Overview", ARTICLE_BLOCKS[1][1]),
                            ("Why it matters", ARTICLE_BLOCKS[2][1]),
                            ("Across the region", ARTICLE_BLOCKS[3][1]))
        self.assertIn("stock_subhead", codes(report(stock)))

    def test_repetitive_caveating_is_flagged_but_one_real_limit_is_not(self):
        heavy = report(legacy_manuscript())
        for code in ("caveat_density", "caveat_endings", "researcher_limits_in_article",
                     "significance_stock"):
            self.assertIn(code, codes(heavy))
        clean = report(article_manuscript())
        self.assertFalse({"caveat_density", "caveat_endings", "caveat_stacking"} & codes(clean))
        # The qualification a reader needs is still in the article text.
        article = " ".join(text for _, text in ARTICLE_BLOCKS)
        self.assertIn("Both figures come from the navy", article)
        self.assertIn("did not name the submarine", article)

    def test_checks_never_edit_or_delete_prose(self):
        before = copy.deepcopy(article_manuscript())
        subject = copy.deepcopy(before)
        report(subject)
        self.assertEqual(subject, before)

    def test_single_country_story_is_not_forced_into_a_comparison(self):
        m = article_manuscript()
        self.assertFalse(any("region" in b["heading"].lower() or "comparison" in b["heading"].lower()
                             for b in m["article"]))
        self.assertIs(validate_manuscript(m, evidence()), m)
        self.assertNotIn("comparison_without_two_desks", codes(report(m)))
        forced = with_blocks(*ARTICLE_BLOCKS[:3], ("Across the region", ARTICLE_BLOCKS[3][1]))
        self.assertIn("comparison_without_two_desks", codes(report(forced)))
        self.assertIn("single-country", m["audit"]["cross_desk_assessment"].lower())

    def test_terms_are_explained_without_changing_the_claim(self):
        plain = with_blocks(("", "The Joint Electronic Warfare Centre (JEWC) opened a new range on 3 "
                                 "March and the JEWC said the first trial would run in June, using "
                                 "equipment the centre built itself in 2025."), *ARTICLE_BLOCKS[1:])
        self.assertNotIn("unexplained_acronym", codes(report(plain)))
        bare = with_blocks(("", "The JEWC opened a new range on 3 March and said the first trial would "
                                "run in June, using equipment the centre built itself in 2025."),
                           *ARTICLE_BLOCKS[1:])
        found = [f for f in report(bare)["flags"] if f["code"] == "unexplained_acronym"]
        self.assertTrue(found and "JEWC" in found[0]["detail"])

    def test_mechanical_errors_are_major_and_set_the_hold(self):
        broken = with_blocks(
            ("", "The navy said the array heard a submarine from 11 kilometres. TODO check this figure "
                 "before the the Brief goes out."), *ARTICLE_BLOCKS[1:3],
            ("Close", "The next step is the autumn exercise, where the figure"))
        result = report(broken)
        self.assertTrue(result["hold"])
        self.assertTrue({"placeholder_text", "truncated_block"} <= codes(result, "errors"))
        self.assertIn("doubled_word", codes(result, "errors"))

    def test_report_text_never_quotes_the_article(self):
        from core.brief_prose_standard import format_report_lines
        text = "\n".join(format_report_lines(report(legacy_manuscript())))
        self.assertNotIn("Marisol", text)
        self.assertTrue(text.endswith("END OF NARRATIVE QUALITY REPORT"))


class ArticleValidationTests(unittest.TestCase):
    """Acceptance checks 8 and 11 for the writer's manuscript validation."""

    def check(self, mutate, pattern):
        m = article_manuscript()
        mutate(m)
        with self.assertRaisesRegex(ValueError, pattern):
            validate_manuscript(m, evidence())

    def test_valid_article_is_accepted_and_ids_are_checked(self):
        m = article_manuscript()
        self.assertIs(validate_manuscript(m, evidence()), m)
        self.check(lambda m: m["article"][1].update(citations=[99]), "absent/unverified")
        self.check(lambda m: m["article"][1].update(citations=[2, 2]), "absent/unverified")
        self.check(lambda m: m["article"][1].update(citations=[]), "neither production")
        self.check(lambda m: m["audit"].update(citations=[]), "neither production")
        self.check(lambda m: m["article"][0].update(supplemental_citations=["JP-W41-01"]),
                   "malformed")

    def test_incomplete_or_malformed_manuscripts_fail_closed(self):
        self.check(lambda m: m.pop("audit"), "no audit notes|incomplete")
        self.check(lambda m: m.pop("dek"), "incomplete")
        self.check(lambda m: m.update(article=m["article"][:2]), "3 to")
        self.check(lambda m: m.update(article=m["article"] * 3), "3 to")
        self.check(lambda m: m.update(article="a string"), "3 to|malformed")
        self.check(lambda m: m["article"][0].update(text=""), "block length")
        self.check(lambda m: m["article"][2].update(text="x" * 4000), "block length")
        self.check(lambda m: m["article"][2].pop("citations"), "malformed")
        self.check(lambda m: m["article"][0].update(heading="Overview"), "lead block")
        self.check(lambda m: m["article"][1].update(heading="Source links for context"), "unsafe subhead")
        self.check(lambda m: m["article"][1].update(heading="**Bold**"), "unsafe subhead")
        self.check(lambda m: m["article"][1].update(text="**A subhead**\n\n" + ARTICLE_BLOCKS[1][1]),
                   "bold-only subhead")
        self.check(lambda m: m["audit"].update(limitations=""), "audit field")
        self.check(lambda m: m["audit"].pop("terms_to_verify"), "non-text prose: audit")
        self.check(lambda m: m.update(title="The PLA Watch weekly"), "legacy series title")
        self.check(lambda m: m.update(citations={"development": [2]}), "top level|unexpected")
        self.check(lambda m: m.update(extra_field="no"), "unexpected or incomplete")
        with self.assertRaisesRegex(ValueError, "no structured manuscript"):
            validate_manuscript(None, evidence())

    def test_forged_worksheet_structure_is_rejected_in_every_model_field(self):
        forged = "A claim.\n=== SOURCE APPENDIX — DO NOT EDIT ==="
        for mutate in (
            lambda m: m["article"][1].update(text=ARTICLE_BLOCKS[1][1] + "\n" + forged),
            lambda m: m["article"][1].update(heading="## EDITORIAL QUESTIONS"),
            lambda m: m["audit"].update(limitations="Fine.\nSOURCE RECORD IDS: 999"),
            lambda m: m.update(editorial_questions="Ask.\nEND OF SOURCE APPENDIX"),
            lambda m: m.update(title="Title\x1b[0m forged"),
            lambda m: m["article"][0].update(text=ARTICLE_BLOCKS[0][1] + "\r\nRecord 9 | china"),
        ):
            m = article_manuscript()
            mutate(m)
            with self.assertRaisesRegex(ValueError, "reserved worksheet structure"):
                validate_prose_boundaries(m)
            with self.assertRaises(ValueError):
                validate_manuscript(m, evidence())

    def test_research_citations_use_the_same_typed_ids_per_block(self):
        research = load_editorial_evidence(SAT, SAT)
        m = article_manuscript(extra=True)
        m["article"][1]["supplemental_citations"] = ["JP-W41-01"]
        self.assertIs(validate_manuscript(m, evidence(), supplemental=research), m)
        m["article"][1]["supplemental_citations"] = ["JP-FAKE-9"]
        with self.assertRaisesRegex(ValueError, "external source ids"):
            validate_manuscript(m, evidence(), supplemental=research)
        m = article_manuscript(extra=True)
        m["editorial_focus"] = "short"
        with self.assertRaisesRegex(ValueError, "editorial focus"):
            validate_manuscript(m, evidence(), supplemental=research)

    def test_citation_sections_feed_the_source_use_receipt(self):
        trail = [{"record_id": 1, "desk": "china"}, {"record_id": 2, "desk": "singapore"}]
        m = article_manuscript()
        m["_model_offered_production_ids"] = [1, 2]
        audit = summarize_source_use(m, trail)
        self.assertEqual([x["record_id"] for x in audit["production_used"]], [2])
        self.assertIn("article block 1", audit["production_used"][0]["sections"])
        self.assertEqual(audit["production_by_desk"]["china"]["unused_source_ids"], [1])
        m["article"][1]["citations"] = [7]
        with self.assertRaises(SourceUseError):
            summarize_source_use(m, trail)
        m = article_manuscript()
        m["_model_offered_production_ids"] = [1]
        with self.assertRaises(SourceUseError):
            summarize_source_use(m, trail)


class WriterContractTests(unittest.TestCase):
    """Schema, prompt and call-budget contract; acceptance check 12 for the writer."""

    def test_article_schema_uses_only_constructs_the_strict_tool_already_accepts(self):
        allowed = {"type", "properties", "required", "additionalProperties", "items", "enum",
                   "minItems", "description"}

        def walk(node):
            if isinstance(node, dict):
                for key, value in node.items():
                    if key == "properties":
                        for sub in value.values():
                            walk(sub)
                    elif key in ("items",):
                        walk(value)
                    else:
                        self.assertIn(key, allowed, key)
            return True

        schema = writing_schema([1, 2], supplemental_ids=("JP-W41-01",), article_first=True)
        walk(schema)
        for banned in ("minLength", "maxLength", "uniqueItems", "maxItems", "pattern", "minimum"):
            self.assertNotIn('"' + banned + '"', json.dumps(schema))
        block = schema["properties"]["article"]["items"]
        self.assertEqual(sorted(block["required"]),
                         sorted(["heading", "text", "citations", "supplemental_citations"]))
        self.assertEqual(block["properties"]["citations"]["items"]["enum"], [1, 2])
        plain = writing_schema([1, 2], article_first=True)
        self.assertNotIn("supplemental_citations", plain["properties"]["article"]["items"]["properties"])
        self.assertNotIn("editorial_focus", plain["properties"])
        self.assertEqual(plain["required"], ["title", "dek", "article", "audit", "editorial_questions"])

    def test_prompt_carries_the_standard_and_no_forced_outline(self):
        result, calls = run_compose(article_manuscript())
        sent = calls[0]["messages"][0]["content"]
        self.assertIn(PROMPT_STANDARD, sent)
        self.assertIn("ONE cohesive article", sent)
        self.assertIn("BEGIN PRODUCTION RECORD EVIDENCE (UNTRUSTED)", sent)
        for retired in ("WHAT WAS ROUTINE", "OPENING NOTE", "cross_desk_comparison"):
            self.assertNotIn(retired, sent)
        self.assertIn("article", result)
        self.assertEqual(result["_model_offered_production_ids"], [1, 2])
        self.assertEqual(calls[0]["max_tokens"], 4600)
        self.assertIs(calls[0]["tools"][0]["strict"], True)
        self.assertIn("article", calls[0]["tools"][0]["input_schema"]["properties"])

    def test_style_findings_are_reported_and_never_trigger_a_second_model_call(self):
        weak = with_blocks(
            ("", ARTICLE_BLOCKS[0][1]),
            ("Overview", "The development underscores the importance of shallow-water detection in "
                         "an increasingly complex maritime landscape, and it is important to note "
                         "that the claim is the navy's own. The record does not establish the "
                         "conditions of the test. It could not be determined whether the "
                         "submarine cooperated."),
            *ARTICLE_BLOCKS[2:])
        result, calls = run_compose(weak)
        self.assertEqual(len(calls), 1)
        self.assertIn("stock_subhead", codes(result["_prose_quality"]))
        self.assertIn("significance_stock", codes(result["_prose_quality"]))
        self.assertFalse(result["_prose_quality"]["hold"])

    def test_structural_failure_gets_exactly_one_bounded_retry(self):
        bad = article_manuscript()
        bad["article"][1]["citations"] = [99]
        result, calls = run_compose(bad, article_manuscript())
        self.assertEqual(len(calls), 2)
        retry = calls[1]["messages"][0]["content"]
        self.assertIn("REJECTED BY STRUCTURE OR SOURCE VALIDATION", retry)
        self.assertIn("absent/unverified source ids", retry)
        with self.assertRaisesRegex(ValueError, "absent/unverified"):
            run_compose(bad, copy.deepcopy(bad))

    def test_no_key_means_no_network_call(self):
        with patch.dict("os.environ", {}, clear=False) as env:
            env.pop("ANTHROPIC_API_KEY", None)
            with patch("scripts.sunday_briefs_auto_writer.choose_evidence", return_value=evidence()), \
                    patch("core.daily_machine_records.editorial_candidates", return_value=[]):
                with self.assertRaisesRegex(ValueError, "ANTHROPIC_API_KEY missing"):
                    compose({"week_start": "2026-10-04", "week_ending": SAT,
                             "desks": ["china"], "source_trail": []}, SAT)

    def test_legacy_shape_is_still_available_on_request(self):
        result, calls = run_compose(legacy_manuscript(), article_first=False)
        self.assertIn("citations", result)
        self.assertNotIn("_prose_quality", result)
        self.assertIn("citations", calls[0]["tools"][0]["input_schema"]["properties"])


class PacketAndProjectionTests(unittest.TestCase):
    """Acceptance checks 4, 8 and 9: article first, evidence layer preserved, no private notes."""

    @classmethod
    def setUpClass(cls):
        cls.manuscript = article_manuscript()
        cls.manuscript["_model_offered_production_ids"] = [1, 2]
        cls.packet = render_packet(scaffold(), manuscript=cls.manuscript, as_of=SAT)
        cls.editable = cls.packet.split("=== EDITABLE MANUSCRIPT ===")[1].split(
            "=== SOURCE APPENDIX")[0]
        cls.appendix = cls.packet.split("=== SOURCE APPENDIX — DO NOT EDIT ===")[1]

    def test_editable_region_is_title_dek_one_article_and_questions(self):
        headings = re.findall(r"^## (.+)$", self.editable, re.MULTILINE)
        self.assertEqual(headings, ["WORKING TITLE", "DEK", "ARTICLE",
                                    "EDITORIAL QUESTIONS / SATURDAY FOLLOW-UP"])
        body = self.editable.split("## ARTICLE")[1].split("## EDITORIAL QUESTIONS")[0]
        self.assertTrue(body.strip().startswith("The Tavaran Navy has spent three weeks"))
        self.assertNotIn("SOURCE RECORD IDS", self.editable)
        self.assertEqual(re.findall(r"^\*\*(.+)\*\*$", body, re.MULTILINE),
                         ["Why shallow water is hard", "A domestic design", "What would settle it"])
        for stock in ("OPENING NOTE", "WHAT WAS ROUTINE", "WHY IT MATTERS", "CROSS-DESK"):
            self.assertNotIn(stock, self.editable)

    def test_research_notes_stay_in_the_private_appendix(self):
        for text in (AUDIT["baseline_and_routine"], AUDIT["limitations"], AUDIT["terms_to_verify"],
                     AUDIT["cross_desk_assessment"], AUDIT["development"]):
            self.assertNotIn(text, self.editable)
            self.assertIn(text, self.appendix)
        self.assertIn("NARRATIVE QUALITY REPORT", self.appendix)
        self.assertNotIn("NARRATIVE QUALITY", self.editable)

    def test_source_accounting_maps_every_block_to_its_verified_ids(self):
        lines = re.findall(r"^Block (\d+) \| fp ([0-9a-f]{12}) \| records (.+) \| external (.+)$",
                           self.appendix, re.MULTILINE)
        self.assertEqual([x[0] for x in lines], ["1", "2", "3", "4"])
        self.assertEqual({x[2] for x in lines}, {"2"})
        for (number, fp, _, _), block in zip(lines, self.manuscript["article"]):
            self.assertEqual(fp, fingerprint(block["text"]))
        self.assertIn("Audit notes | records 2 | external none", self.appendix)
        self.assertIn("- singapore: 1 in human appendix; 1 actually model-offered; 1 cited.",
                      self.appendix)
        self.assertIn("- china: 1 in human appendix; 1 actually model-offered; 0 cited.",
                      self.appendix)

    def test_header_stays_unnumbered_and_unapproved(self):
        self.assertIn("Status: UNNUMBERED DRAFT — NOT APPROVED OR PUBLISHED", self.packet)
        self.assertIn("You do not supply record IDs or settle permissions", self.packet)
        self.assertNotIn("Give record IDs for factual claims", self.packet)
        self.assertTrue(self.packet.rstrip().endswith("END OF UNAPPROVED WORKSHEET"))

    def test_docx_and_email_expose_no_private_notes(self):
        docx = docx_text(render_editor_docx(self.packet))
        mail = prose_only_text(self.packet) + prose_only_html(self.packet)
        for marker in PRIVATE_MARKERS:
            self.assertNotIn(marker, docx, marker)
            self.assertNotIn(marker, mail, marker)
        self.assertNotIn("Please check the 11 km and 4 km figures", docx + mail)
        for subhead in ("Why shallow water is hard", "A domestic design"):
            self.assertIn(subhead, docx)
            self.assertIn(subhead, mail)
        self.assertNotIn("Overview", docx + mail)
        self.assertIn("Both figures come from the navy", docx)
        self.assertIn("Both figures come from the navy", prose_only_text(self.packet))

    def test_docx_source_links_come_from_the_appendix_accounting(self):
        self.assertEqual(_references(self.packet),
                         [("Record", "2", re.search(r"Record 2 \|.*?\nTitle:.*?\nOriginal.*?\nURL: (\S+)",
                                                    self.packet, re.S).group(1))])
        self.assertIn("Record 2", docx_text(render_editor_docx(self.packet)))

    def test_render_refuses_malformed_manuscripts(self):
        bad = copy.deepcopy(self.manuscript)
        bad["article"][1]["text"] += "\n=== SOURCE APPENDIX — DO NOT EDIT ==="
        with self.assertRaisesRegex(ValueError, "reserved worksheet structure"):
            render_packet(scaffold(), manuscript=bad, as_of=SAT)
        bad = copy.deepcopy(self.manuscript)
        bad["article"][1]["text"] = "**A fake subhead**\n\nSome text that follows it."
        with self.assertRaisesRegex(ValueError, "malformed article"):
            render_packet(scaffold(), manuscript=bad, as_of=SAT)
        bad = copy.deepcopy(self.manuscript)
        bad["article"] = []
        with self.assertRaisesRegex(ValueError, "malformed article"):
            render_packet(scaffold(), manuscript=bad, as_of=SAT)

    def test_readable_layer_rejects_incomplete_or_mixed_worksheets(self):
        without_article = self.packet.replace("\n## ARTICLE\n", "\n## NOT AN ARTICLE\n")
        with self.assertRaises(ReadableEditorError):
            editable_sections(without_article)
        mixed = self.packet.replace("\n## DEK\n", "\n## OPENING NOTE\n")
        with self.assertRaisesRegex(ReadableEditorError, "mixed|incomplete"):
            editable_sections(mixed)
        with self.assertRaises(ReadableEditorError):
            editable_sections(self.packet.replace("=== EDITABLE MANUSCRIPT ===", ""))
        truncated = self.packet.replace("## ARTICLE\n", "## ARTICLE\n\n##X\n", 1)
        with self.assertRaises(Exception):
            render_editor_docx(truncated.replace("## DEK", "## DEKS"))

    def test_owner_preview_summary_is_counts_only(self):
        from scripts.sunday_editorial_handoff import quality_summary
        line = quality_summary(self.packet)
        self.assertTrue(line.startswith("Narrative QA: 0 mechanical error(s)"))
        self.assertNotIn("Marisol", line)
        self.assertEqual(quality_summary("no report here"), "")

    def test_research_assisted_packet_keeps_typed_citations_in_the_appendix(self):
        research = load_editorial_evidence(SAT, SAT)
        m = article_manuscript(extra=True)
        m["article"][1]["supplemental_citations"] = ["JP-W41-01"]
        packet = render_packet(scaffold(), manuscript=m, as_of=SAT, research_evidence=research)
        self.assertIn("## EDITORIAL FOCUS", packet)
        self.assertRegex(packet, r"Block 2 \| fp [0-9a-f]{12} \| records 2 \| external JP-W41-01")
        self.assertNotIn("EXTERNAL SOURCE IDS", packet.split("=== SOURCE APPENDIX")[0])
        self.assertIn("External source JP-W41-01 | japan", packet)
        self.assertEqual(_references(packet)[-1][:2], ("External", "JP-W41-01"))
        self.assertNotIn("JP-W41-01", docx_text(render_editor_docx(packet)).split(
            "Source links for context")[0])


class ReturnWorkflowTests(unittest.TestCase):
    """Acceptance check 10: Dylan's return validation and Word workflow."""

    @classmethod
    def setUpClass(cls):
        cls.packet = render_packet(scaffold(), manuscript=article_manuscript(), as_of=SAT)

    def test_reordered_merged_and_rewritten_paragraphs_validate(self):
        lead, second = ARTICLE_BLOCKS[0][1].split("\n\n")
        edited = self.packet.replace(
            lead + "\n\n" + second,
            "Weeks of towing ended on 14 October, when the Tavaran Navy said its new array had heard a "
            "submarine from 11 kilometres away in shallow water.\n\n" + second)
        edited = edited.replace("**A domestic design**\n\n", "")
        receipt = validate_return(self.packet, edited)
        self.assertEqual(receipt["citation_lines"], 0)
        self.assertEqual(receipt["review_status"], "STRUCTURAL REVIEW ONLY — UNAPPROVED")
        paras = ARTICLE_BLOCKS[3][1].split("\n\n")
        swapped = self.packet.replace("\n\n".join(paras), "\n\n".join(reversed(paras)))
        self.assertGreater(validate_return(self.packet, swapped)["changed_sections"], 0)

    def test_return_that_touches_protected_parts_is_refused(self):
        edited = self.packet.replace("The Tavaran Navy has spent", "The navy spent", 1)
        validate_return(self.packet, edited)
        with self.assertRaisesRegex(ReturnValidationError, "no edits"):
            validate_return(self.packet, self.packet)
        for protected in (AUDIT["limitations"][:30], "Block 2 | fp ", "records 2 | external none"):
            with self.subTest(protected=protected):
                with self.assertRaisesRegex(ReturnValidationError, "appendix"):
                    validate_return(self.packet, edited.replace(protected, "X" + protected, 1))
        with self.assertRaisesRegex(ReturnValidationError, "heading"):
            validate_return(self.packet, edited.replace("\n## ARTICLE\n", "\n## BODY\n"))
        with self.assertRaisesRegex(ReturnValidationError, "no prose|blank"):
            validate_return(self.packet, re.sub(r"(## ARTICLE\n).*?(\n## EDITORIAL)", r"\1\2",
                                                 edited, flags=re.S))
        with self.assertRaisesRegex(ReturnValidationError, "citation"):
            validate_return(self.packet, edited.replace("**Why shallow water is hard**",
                                                        "SOURCE RECORD IDS: 2", 1))

    def test_word_round_trip_keeps_source_accounting_and_accepts_subhead_edits(self):
        data = render_editor_docx(self.packet)

        def mutate(doc):
            paragraphs = doc.paragraphs
            lead = next(p for p in paragraphs if p.text.startswith("The Tavaran Navy has spent"))
            set_text(lead, "For three weeks the navy towed a new sonar array in the Kesh Strait.")
            heading = next(p for p in paragraphs if p.text == "A domestic design")
            set_text(heading, "Built at home")
            gone = next(p for p in paragraphs if p.text == "What would settle it")
            gone._element.getparent().remove(gone._element)

        candidate, receipt = reconstruct_return(self.packet, edit_docx(data, mutate))
        self.assertEqual(receipt["review_status"], "STRUCTURAL REVIEW ONLY — UNAPPROVED")
        header_and_appendix = lambda t: (t.split("=== EDITABLE MANUSCRIPT ===")[0],
                                         t.split("=== SOURCE APPENDIX — DO NOT EDIT ===")[1])
        self.assertEqual(header_and_appendix(candidate)[1].rstrip("\n"),
                         header_and_appendix(self.packet)[1].rstrip("\n"))
        article = candidate.split("## ARTICLE")[1].split("## EDITORIAL QUESTIONS")[0]
        self.assertIn("**Built at home**", article)
        self.assertNotIn("What would settle it", article)
        self.assertIn("## EDITORIAL QUESTIONS / SATURDAY FOLLOW-UP", candidate)
        self.assertIn("Please check the 11 km and 4 km figures", candidate)

    def test_word_return_fails_safely(self):
        data = render_editor_docx(self.packet)
        with self.assertRaisesRegex(ReturnValidationError, "no edits"):
            reconstruct_return(self.packet, data)
        with self.assertRaises(ReturnValidationError):
            reconstruct_return(self.packet, b"not a docx")
        with self.assertRaises(ReturnValidationError):
            reconstruct_return(self.packet, b"")

        def remove_title(doc):
            title = next(p for p in doc.paragraphs if p.style.name == "Title")
            title._element.getparent().remove(title._element)

        with self.assertRaises(ReturnValidationError):
            reconstruct_return(self.packet, edit_docx(data, remove_title))

        def forge(doc):
            lead = next(p for p in doc.paragraphs if p.text.startswith("The Tavaran Navy has spent"))
            set_text(lead, "SOURCE RECORD IDS: 2 and then some prose that is long enough.")

        with self.assertRaisesRegex(ReturnValidationError, "protected source markers"):
            reconstruct_return(self.packet, edit_docx(data, forge))

        def early_subhead(doc):
            title = next(p for p in doc.paragraphs if p.style.name == "Title")
            title.insert_paragraph_before("Early heading", style="Heading 2")

        with self.assertRaisesRegex(ReturnValidationError, "heading"):
            reconstruct_return(self.packet, edit_docx(data, early_subhead))

        def tracked(doc):
            lead = next(p for p in doc.paragraphs if p.text.startswith("The Tavaran Navy has spent"))
            from docx.oxml import OxmlElement
            lead._p.append(OxmlElement("w:ins"))

        with self.assertRaisesRegex(ReturnValidationError, "tracked changes"):
            reconstruct_return(self.packet, edit_docx(data, tracked))


class LegacyCompatibilityTests(unittest.TestCase):
    """Worksheets already sent in the fixed-outline layout must keep working."""

    def test_legacy_packet_layout_and_word_round_trip_are_unchanged(self):
        legacy = render_packet(scaffold(), manuscript=legacy_manuscript(), as_of=SAT)
        headings = re.findall(r"^## (.+)$", legacy, re.MULTILINE)
        self.assertEqual(headings[:3], ["WORKING TITLE", "DEK", "CONCRETE DEVELOPMENT"])
        self.assertIn("SOURCE RECORD IDS: 2", legacy)
        self.assertNotIn("ARTICLE SOURCE ACCOUNTING", legacy)
        self.assertIn("Give record IDs for factual claims", legacy)
        data = render_editor_docx(legacy)
        self.assertIn("Overview", docx_text(data))

        def mutate(doc):
            lead = next(p for p in doc.paragraphs if p.text.startswith("On 16 October"))
            set_text(lead, "On 14 October the Tavaran Navy described a sonar trial on the patrol ship Marisol.")

        candidate, receipt = reconstruct_return(legacy, edit_docx(data, mutate))
        self.assertEqual(receipt["citation_lines"], 7)
        self.assertEqual(len(_references(legacy)), 1)


class EditorReturnAnalysisTests(unittest.TestCase):
    """The offline feedback loop: observations only, never ingestion or promotion."""

    @classmethod
    def setUpClass(cls):
        cls.packet = render_packet(scaffold(), manuscript=article_manuscript(), as_of=SAT)
        lead = ARTICLE_BLOCKS[0][1].split("\n\n")[0]
        cut = ARTICLE_BLOCKS[1][1].split("\n\n")[2]
        edited = cls.packet.replace(lead, "For three weeks the navy towed a sonar array in the Kesh "
                                          "Strait, and on 14 October it said what it found.")
        edited = edited.replace("\n\n" + cut, "\n\nThe array reportedly cost 40 million.")
        cls.edited = edited
        cls.report = analyzer.analyze(cls.packet, edited)

    def test_patterns_cited_blocks_and_figures_are_reported(self):
        r = self.report
        self.assertEqual(r["layout"], "article-first")
        # The lead paragraph was replaced wholesale by a much shorter one (cut + added)
        # and block 2 lost a paragraph to a one-line insertion (cut + added).
        self.assertEqual(r["counts"]["cut_or_merged"], 2)
        self.assertEqual(r["counts"]["added"], 2)
        verdicts = {b["block"]: b["verdict"] for b in r["cited_blocks"]}
        self.assertEqual(verdicts[3], "intact")
        self.assertIn("cut or merged", verdicts[1])
        self.assertIn("cut or merged", verdicts[2])
        self.assertEqual(r["cited_blocks"][1]["records"], ["2"])
        self.assertIn("40", r["figures"]["new_in_returned"])
        self.assertIn("cut_paragraphs", {p["code"] for p in r["candidate_patterns"]})
        self.assertFalse(any(v for k, v in r["governance"].items() if k.endswith(("changed", "called", "ingested_correspondence"))))
        self.assertIn("never changes the prompt", r["governance"]["promotion_rule"])
        markdown = analyzer.render_markdown(r)
        self.assertIn("Block 2 (records 2; external none)", markdown)
        self.assertIn("No correspondence was read", markdown)

    def test_a_pure_reorder_is_seen_as_movement_not_rewriting(self):
        paras = ARTICLE_BLOCKS[3][1].split("\n\n")
        edited = self.packet.replace("\n\n".join(paras), "\n\n".join(reversed(paras)))
        r = analyzer.analyze(self.packet, edited)
        self.assertEqual((r["counts"]["rewritten"], r["counts"]["cut_or_merged"]), (0, 0))
        self.assertGreaterEqual(r["counts"]["moved"], 1)
        self.assertIn("moved", {b["block"]: b["verdict"] for b in r["cited_blocks"]}[4])

    def test_analysis_reuses_return_validation_and_refuses_invalid_returns(self):
        with self.assertRaises(ReturnValidationError):
            analyzer.analyze(self.packet, self.packet)
        with self.assertRaises(ReturnValidationError):
            analyzer.analyze(self.packet, self.edited.replace("=== SOURCE APPENDIX", "=== X APPENDIX"))

    def test_legacy_returns_can_be_analysed_too(self):
        legacy = render_packet(scaffold(), manuscript=legacy_manuscript(), as_of=SAT)
        edited = legacy.replace("I am watching whether the navy publishes further information. "
                                "I could not determine when this might occur.",
                                "The navy has promised no date for a fuller release.")
        r = analyzer.analyze(legacy, edited)
        self.assertEqual(r["layout"], "fixed-outline")
        self.assertEqual(r["cited_blocks"], [])
        self.assertIn("removed_caveat_formulas", {p["code"] for p in r["candidate_patterns"]})

    def test_aggregation_never_promotes_a_rule(self):
        rows = analyzer.aggregate([self.report])
        self.assertTrue(all(r["status"] == "isolated observation" and r["promoted"] is False for r in rows))
        rows = analyzer.aggregate([self.report] * 3)
        self.assertTrue(all("owner consideration" in r["status"] and r["promoted"] is False for r in rows))
        with self.assertRaises(ValueError):
            analyzer.aggregate([{"schema": "other"}])

    def test_cli_writes_private_files_outside_the_repository_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "sent.txt").write_text(self.packet, encoding="utf-8")
            (d / "back.txt").write_text(self.edited, encoding="utf-8")
            out, js = d / "report.md", d / "report.json"
            argv = ["--original", str(d / "sent.txt"), "--edited", str(d / "back.txt"),
                    "--out", str(out), "--json-out", str(js)]
            self.assertEqual(analyzer.main(argv), 0)
            self.assertIn("Candidate patterns", out.read_text(encoding="utf-8"))
            self.assertEqual(json.loads(js.read_text())["schema"], analyzer.SCHEMA)
            self.assertEqual(oct(out.stat().st_mode & 0o777), "0o600")
            with self.assertRaises(SystemExit):
                analyzer.main(argv)  # never overwrites
            inside = analyzer.REPO_ROOT / "tmp-analysis-should-not-exist.md"
            with self.assertRaises(SystemExit):
                analyzer.main(["--original", str(d / "sent.txt"), "--edited", str(d / "back.txt"),
                               "--out", str(inside)])
            self.assertFalse(inside.exists())
            word = d / "back.docx"
            word.write_bytes(edit_docx(render_editor_docx(self.packet), lambda doc: set_text(
                next(p for p in doc.paragraphs if p.text.startswith("The Tavaran Navy has spent")),
                "For three weeks the navy towed a new sonar array in the Kesh Strait.")))
            self.assertEqual(analyzer.main(["--original", str(d / "sent.txt"), "--edited-word", str(word),
                                            "--json-out", str(d / "word.json")]), 0)
            self.assertEqual(analyzer.main(["--aggregate", str(js), str(d / "word.json")]), 0)

    def test_analyzer_has_no_mail_or_model_dependency(self):
        source = Path(analyzer.__file__).read_text(encoding="utf-8")
        for forbidden in ("smtplib", "imaplib", "anthropic", "requests", "urllib", "sqlite3",
                          "subprocess", "git "):
            self.assertNotIn("import " + forbidden.strip(), source)


class NoSideEffectTests(unittest.TestCase):
    """Acceptance check 12: rendering, projecting and analysing never send or publish."""

    def test_no_smtp_no_numbering_no_publication_in_the_offline_workflow(self):
        with patch("smtplib.SMTP_SSL", side_effect=AssertionError("SMTP used")), \
                patch("smtplib.SMTP", side_effect=AssertionError("SMTP used")):
            packet = render_packet(scaffold(), manuscript=article_manuscript(), as_of=SAT)
            prose_only_text(packet)
            prose_only_html(packet)
            render_editor_docx(packet)
            analyzer.analyze(packet, packet.replace("The Tavaran Navy has spent", "The navy spent", 1))
        self.assertIn("UNNUMBERED DRAFT", packet)
        self.assertNotRegex(packet, r"issue_number|Issue No\.|editorial_status: approved")
        self.assertNotIn("ANTHROPIC_API_KEY", packet)


if __name__ == "__main__":
    unittest.main()
